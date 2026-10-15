import hashlib
import hmac
import json
import time
from datetime import date, datetime, timedelta
from io import BytesIO
from unittest.mock import patch
from urllib.parse import urlencode

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group, Permission
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from servizio.models import ScheduledTask, Servizio, Timbratura
from tg_bot.models import TelegramUser
from volontario.models import Volontario

from .reports import hours_report, month_range
from .telegram_auth import validate_init_data
from .views import XLSX_CONTENT_TYPE, _parse_range

BOT_TOKEN = "123456:TEST-TOKEN"
VALID_CF = "RSSMRA90A01H501W"
OTHER_CF = "VRDLGU85M01F205P"

User = get_user_model()


def sign_init_data(fields: dict, bot_token: str = BOT_TOKEN) -> str:
    """Build an initData query string signed the way Telegram does."""
    data_check_string = "\n".join(f"{k}={v}" for k, v in sorted(fields.items()))
    secret = hmac.new(b"WebAppData", bot_token.encode(), hashlib.sha256).digest()
    digest = hmac.new(secret, data_check_string.encode(), hashlib.sha256).hexdigest()
    return urlencode({**fields, "hash": digest})


def init_data_for(telegram_id: int, auth_date: int | None = None) -> str:
    return sign_init_data(
        {
            "auth_date": str(auth_date or int(time.time())),
            "query_id": "AAH",
            "user": json.dumps({"id": telegram_id, "first_name": "Mario"}),
        }
    )


def aware(*args) -> datetime:
    return timezone.make_aware(datetime(*args))


class ValidateInitDataTests(TestCase):
    def test_valid(self):
        user = validate_init_data(init_data_for(42), BOT_TOKEN)
        self.assertEqual(user["id"], 42)

    def test_wrong_token(self):
        self.assertIsNone(validate_init_data(init_data_for(42), "other:token"))

    def test_tampered(self):
        tampered = init_data_for(42).replace("42", "43")
        self.assertIsNone(validate_init_data(tampered, BOT_TOKEN))

    def test_expired(self):
        old = int(time.time()) - 2 * 24 * 3600
        self.assertIsNone(validate_init_data(init_data_for(42, old), BOT_TOKEN))

    def test_missing_hash_or_garbage(self):
        self.assertIsNone(validate_init_data("user=%7B%7D&auth_date=1", BOT_TOKEN))
        self.assertIsNone(validate_init_data("not a query", BOT_TOKEN))
        self.assertIsNone(validate_init_data("", BOT_TOKEN))


class WebappTestCase(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="mario")
        self.volontario = Volontario.objects.create(
            codice_fiscale=VALID_CF, nome="Mario", cognome="Rossi", user=self.user
        )
        TelegramUser.objects.create(telegram_id=42, volontario=self.volontario)


@override_settings(TELEGRAM_BOT_TOKEN=BOT_TOKEN)
class TelegramAuthViewTests(WebappTestCase):
    url = reverse("webapp:telegram_auth")

    def test_logs_in_linked_user(self):
        response = self.client.post(self.url, {"init_data": init_data_for(42)})
        self.assertEqual(response.json(), {"ok": True})
        self.assertEqual(self.client.get(reverse("webapp:home")).status_code, 200)

    def test_rejects_invalid_signature(self):
        response = self.client.post(self.url, {"init_data": "auth_date=1&hash=abc"})
        self.assertEqual(response.status_code, 403)

    def test_rejects_unlinked_telegram_user(self):
        response = self.client.post(self.url, {"init_data": init_data_for(99)})
        self.assertEqual(response.status_code, 403)
        self.assertFalse(response.json()["ok"])


class AccessTests(WebappTestCase):
    def test_anonymous_gets_gate(self):
        response = self.client.get(reverse("webapp:home"))
        self.assertEqual(response.status_code, 401)
        self.assertTemplateUsed(response, "webapp/gate.html")
        # Embeddable inside Telegram Web's iframe.
        self.assertNotIn("X-Frame-Options", response.headers)

    def test_home_for_volunteer(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("webapp:home"))
        self.assertContains(response, "Ciao Mario")
        self.assertNotContains(response, "Report ore")

    def test_report_forbidden_without_permission(self):
        # Being staff is not enough: the report needs its own permission.
        self.user.is_staff = True
        self.user.save()
        self.client.force_login(self.user)
        response = self.client.get(reverse("webapp:report_ore"))
        self.assertEqual(response.status_code, 403)
        home = self.client.get(reverse("webapp:home"))
        self.assertNotContains(home, reverse("webapp:report_ore"))

    def test_report_with_permission(self):
        self.user.user_permissions.add(
            Permission.objects.get(
                codename="view_report_ore", content_type__app_label="servizio"
            )
        )
        self.client.force_login(self.user)
        home = self.client.get(reverse("webapp:home"))
        self.assertContains(home, reverse("webapp:report_ore"))
        self.assertNotContains(home, "Admin</a>")  # admin link still needs is_staff
        Timbratura.objects.create(
            fkvolontario=self.volontario,
            clock_in=timezone.now() - timedelta(hours=3),
            clock_out=timezone.now() - timedelta(hours=1),
        )
        response = self.client.get(reverse("webapp:report_ore"))
        self.assertContains(response, "Mario Rossi")
        self.assertContains(response, "2h")

        csv_response = self.client.get(reverse("webapp:report_ore"), {"formato": "csv"})
        self.assertEqual(csv_response["Content-Type"], "text/csv; charset=utf-8")
        self.assertIn("Mario Rossi;2,00;1", csv_response.content.decode())

        xlsx_response = self.client.get(
            reverse("webapp:report_ore"), {"formato": "xlsx"}
        )
        self.assertEqual(xlsx_response["Content-Type"], XLSX_CONTENT_TYPE)
        rows = list(
            load_workbook(BytesIO(xlsx_response.content)).active.iter_rows(
                values_only=True
            )
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][:3], (VALID_CF, "Rossi", "Mario"))
        self.assertEqual(rows[1][5], 2)


@patch("servizio.timbrature.invia_messaggi_servizio_terminato_background")
class TerminaServizioTests(WebappTestCase):
    def setUp(self):
        super().setUp()
        now = timezone.now()
        self.servizio = Servizio.objects.create(
            nome="Sagra",
            data_ora=now - timedelta(hours=3),
            data_ora_fine=now + timedelta(hours=2),
            send_message=False,
        )
        self.altro = Servizio.objects.create(
            nome="Fiera", data_ora=now - timedelta(hours=1), send_message=False
        )
        luigi = Volontario.objects.create(
            codice_fiscale=OTHER_CF, nome="Luigi", cognome="Verdi"
        )
        self.aperte = [
            Timbratura.objects.create(
                fkvolontario=vol,
                fkservizio=self.servizio,
                clock_in=now - timedelta(hours=2),
            )
            for vol in (self.volontario, luigi)
        ]
        self.altro_turno = Timbratura.objects.create(
            fkvolontario=Volontario.objects.create(
                codice_fiscale="BNCGNN80A01F205W", nome="Gianni", cognome="Bianchi"
            ),
            fkservizio=self.altro,
            clock_in=now - timedelta(minutes=30),
        )
        self.url = reverse("webapp:termina_servizio", args=[self.servizio.pk])
        self.client.force_login(self.user)

    def make_coordinatore(self):
        self.user.groups.add(Group.objects.create(name="Coordinatore"))

    def test_coordinatore_sees_open_servizi_and_volunteers(self, notify):
        self.make_coordinatore()
        response = self.client.get(reverse("webapp:home"))
        self.assertContains(response, "Termina servizio", count=2)
        self.assertContains(response, "Luigi Verdi")
        self.assertContains(response, "Gianni Bianchi")

    def test_termina_servizio_clocks_out_its_volunteers(self, notify):
        self.make_coordinatore()
        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(self.url)
        self.assertRedirects(response, reverse("webapp:home"))

        for t in self.aperte:
            t.refresh_from_db()
            self.assertIsNotNone(t.clock_out)
        self.altro_turno.refresh_from_db()
        self.assertIsNone(self.altro_turno.clock_out)
        self.servizio.refresh_from_db()
        self.assertEqual(self.servizio.data_ora_fine, self.aperte[0].clock_out)
        notify.assert_called_once_with([t.pk for t in self.aperte])

    def test_volunteer_cannot_terminate(self, notify):
        response = self.client.get(reverse("webapp:home"))
        self.assertNotContains(response, "Termina servizio")
        self.assertEqual(self.client.post(self.url).status_code, 403)
        self.assertFalse(Timbratura.objects.filter(clock_out__isnull=False).exists())


class HoursReportTests(TestCase):
    def setUp(self):
        self.mario = Volontario.objects.create(
            codice_fiscale=VALID_CF, nome="Mario", cognome="Rossi"
        )
        self.luigi = Volontario.objects.create(
            codice_fiscale=OTHER_CF, nome="Luigi", cognome="Verdi"
        )
        servizio = Servizio.objects.create(
            nome="Fiera", data_ora=aware(2026, 3, 1, 9), send_message=False
        )
        task = ScheduledTask.objects.create(
            nome="Controllo mezzi", deadline=aware(2026, 4, 30, 18)
        )
        entries = [
            (self.mario, aware(2026, 3, 1, 9), aware(2026, 3, 1, 13), servizio, None),
            (self.mario, aware(2026, 4, 2, 18), aware(2026, 4, 2, 19, 30), None, task),
            (self.luigi, aware(2026, 4, 5, 10), aware(2026, 4, 5, 11), None, None),
            # Outside the range.
            (self.luigi, aware(2026, 6, 1, 10), aware(2026, 6, 1, 20), None, None),
        ]
        for vol, cin, cout, srv, tsk in entries:
            Timbratura.objects.create(
                fkvolontario=vol,
                clock_in=cin,
                clock_out=cout,
                fkservizio=srv,
                fkscheduled_task=tsk,
            )
        # Still open: counted as open, not as hours.
        Timbratura.objects.create(fkvolontario=self.luigi, clock_in=aware(2026, 4, 9))

    def test_totals_and_breakdowns(self):
        report = hours_report(date(2026, 3, 1), date(2026, 5, 31))

        self.assertAlmostEqual(report.ore_totali, 6.5)
        self.assertEqual(report.turni, 3)
        self.assertEqual(report.volontari_attivi, 2)
        self.assertEqual(report.turni_aperti, 1)

        self.assertEqual(
            [m["label"] for m in report.mesi], ["mar 26", "apr 26", "mag 26"]
        )
        self.assertEqual(report.mesi[0]["servizio"], 4)
        self.assertEqual(report.mesi[1]["attivita"], 1.5)
        self.assertEqual(report.mesi[1]["altro"], 1)
        self.assertEqual(report.mesi[2]["servizio"], 0)

        mario, luigi = report.per_volontario
        self.assertEqual(mario.nome, "Mario Rossi")
        self.assertAlmostEqual(mario.ore, 5.5)
        self.assertEqual(mario.turni, 2)
        self.assertAlmostEqual(luigi.ore, 1)

    def test_end_date_is_inclusive(self):
        report = hours_report(date(2026, 4, 5), date(2026, 4, 5))
        self.assertAlmostEqual(report.ore_totali, 1)

    def test_month_range_crosses_year(self):
        self.assertEqual(
            month_range(date(2025, 11, 15), date(2026, 2, 1)),
            [(2025, 11), (2025, 12), (2026, 1), (2026, 2)],
        )


class ParseRangeTests(TestCase):
    today = date(2026, 10, 9)

    def test_default_last_12_months(self):
        self.assertEqual(
            _parse_range({}, self.today), (date(2025, 11, 1), self.today, "12m")
        )

    def test_presets(self):
        self.assertEqual(
            _parse_range({"periodo": "mese"}, self.today)[0], date(2026, 10, 1)
        )
        self.assertEqual(
            _parse_range({"periodo": "3m"}, self.today)[0], date(2026, 8, 1)
        )
        self.assertEqual(
            _parse_range({"periodo": "anno"}, self.today)[0], date(2026, 1, 1)
        )

    def test_custom_range_and_invalid(self):
        self.assertEqual(
            _parse_range({"da": "2026-01-10", "a": "2026-02-20"}, self.today),
            (date(2026, 1, 10), date(2026, 2, 20), ""),
        )
        # Reversed or malformed dates fall back to the default.
        self.assertEqual(
            _parse_range({"da": "2026-03-01", "a": "2026-01-01"}, self.today)[2], "12m"
        )
        self.assertEqual(_parse_range({"da": "x", "a": "y"}, self.today)[2], "12m")
